"""Read-only source extraction with page/slide/cell provenance for this review."""
import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import pdfplumber

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[3]
WORK = Path(__file__).resolve().parent
RESOURCES = ROOT.parents[2] / '50-资源'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
PAPERS = [
    ('P01', '高效跨范式类脑计算芯片'),
    ('P02', '神经形态中间表示'),
    ('P03', '硕士论文_曲环宇'),
    ('P04', '基于混合型Tianjic'),
    ('P05', '分布众核类脑芯片'),
    ('P06', 'Tianjic：'),
    ('P07', 'TianjicX1概要设计'),
    ('P08', 'SNN - MLIR'),
    ('P09', 'MLIR_Scaling'),
    ('P10', 'BIVM：'),
]


def dump(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding='utf-8')


def extract_pdf(path, out):
    pages = []
    sparse = []
    engine = 'pypdf'
    try:
        reader = PdfReader(path)
        texts = [page.extract_text() or '' for page in reader.pages]
    except Exception:
        engine = 'pdfplumber'
        with pdfplumber.open(path) as pdf:
            texts = [page.extract_text() or '' for page in pdf.pages]
    for number, content in enumerate(texts, 1):
        content = content.encode('utf-8', errors='replace').decode('utf-8')
        if len(content.strip()) < 80:
            sparse.append(number)
        pages.append(f'\n## PDF第{number}页\n\n{content}\n')
        dump(out / f'页{number:03d}.md', pages[-1])
    dump(out / '全文.md', f'# {path.name}\n\n来源：{path}\n' + ''.join(pages))
    for start in range(0, len(pages), 10):
        dump(out / f'页{start+1:03d}-{min(start+10,len(pages)):03d}.md', ''.join(pages[start:start+10]))
    return {'pages': len(pages), 'sparse_pages': sparse,
            'text_chars': sum(len(p) for p in pages),
            'engine': engine}


def extract_docx(path, out):
    doc = Document(path)
    parts = [f'# {path.name}\n\n来源：{path}\n']
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            parts.append(f'段落{i} [{para.style.name}]：{para.text}\n')
    for i, table in enumerate(doc.tables, 1):
        parts.append(f'\n## 表{i}\n')
        for r, row in enumerate(table.rows, 1):
            parts.append(f'行{r}：' + ' | '.join(cell.text for cell in row.cells) + '\n')
    dump(out / '全文.md', '\n'.join(parts))
    return {'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}


def extract_pptx(path, out):
    parts = [f'# {path.name}\n\n来源：{path}\n']
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                        key=lambda n: int(re.search(r'(\d+)\.xml', n).group(1)))
        image_counts = []
        for i, name in enumerate(slides, 1):
            node = ET.fromstring(z.read(name))
            texts = [e.text or '' for e in node.findall('.//a:t', NS)]
            # Preserve shape boundaries and connector metadata; not a rendered slide.
            shapes = []
            for element in node.iter():
                tag = element.tag.rsplit('}', 1)[-1]
                if tag not in ('sp', 'cxnSp', 'pic'):
                    continue
                shape_text = '\n'.join(e.text or '' for e in element.findall('.//a:t', NS))
                ident = next((e.attrib for e in element.iter()
                              if e.tag.rsplit('}', 1)[-1] == 'cNvPr'), {})
                geometry = [{'tag': e.tag.rsplit('}', 1)[-1], **e.attrib}
                            for e in element.iter()
                            if e.tag.rsplit('}', 1)[-1] in
                            ('off', 'ext', 'stCxn', 'endCxn', 'headEnd', 'tailEnd', 'prstGeom')]
                shapes.append({'type': tag, 'identity': ident,
                               'text': shape_text, 'geometry': geometry})
            dump(out / f'幻灯片{i:03d}_结构.json',
                 json.dumps(shapes, ensure_ascii=False, indent=2))
            image_count = len(node.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}pic'))
            image_counts.append(image_count)
            part = f'\n## 幻灯片{i}\n\n图片对象：{image_count}\n\n' + '\n'.join(texts)
            notes = f'ppt/notesSlides/notesSlide{i}.xml'
            if notes in z.namelist():
                note = ET.fromstring(z.read(notes))
                part += '\n\n备注：\n' + '\n'.join(e.text or '' for e in note.findall('.//a:t', NS))
            parts.append(part)
            dump(out / f'幻灯片{i:03d}.md', part)
        media = [n for n in z.namelist() if n.startswith('ppt/media/') and not n.endswith('/')]
        for name in media:
            target = out / '媒体' / Path(name).name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(z.read(name))
    dump(out / '全文.md', '\n'.join(parts))
    return {'slides': len(slides), 'image_counts': image_counts, 'media': len(media),
            'limitation': 'Text and media extraction does not verify diagram connections or visual layout.'}


def extract_xlsx(path, out):
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    parts = [f'# {path.name}\n\n来源：{path}\n']
    sheets = []
    for sheet in formulas:
        part = [f'\n## 工作表 {sheet.title}\n\n范围：{sheet.calculate_dimension()}\n合并：{list(sheet.merged_cells.ranges)}\n']
        count = 0
        styled = []
        for row in sheet:
            for cell in row:
                if cell.value is not None or cell.has_style:
                    styled.append({'cell': cell.coordinate, 'value': str(cell.value) if cell.value is not None else None,
                                   'style_id': cell.style_id, 'fill': cell.fill.patternType,
                                   'fgColor': str(cell.fill.fgColor),
                                   'border': [getattr(cell.border, side).style if getattr(cell.border, side) else None
                                              for side in ('left', 'right', 'top', 'bottom')]})
                if cell.value is not None:
                    cached = values[sheet.title][cell.coordinate].value
                    part.append(f'{cell.coordinate}：{cell.value}' + (f' [缓存值：{cached}]' if cell.data_type == 'f' else ''))
                    count += 1
        text = '\n'.join(part)
        parts.append(text)
        dump(out / (re.sub(r'[\\/:*?"<>|]', '_', sheet.title) + '.md'), text)
        dump(out / (re.sub(r'[\\/:*?"<>|]', '_', sheet.title) + '_结构.json'),
             json.dumps({'cells': styled, 'merged': [str(r) for r in sheet.merged_cells.ranges]},
                        ensure_ascii=False, indent=2))
        sheets.append({'name': sheet.title, 'range': sheet.calculate_dimension(), 'nonempty_cells': count})
    dump(out / '全文.md', '\n'.join(parts))
    formulas.close()
    values.close()
    return {'sheets': sheets}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--only', default='all')
    args = parser.parse_args()
    items = []
    for ident, prefix in PAPERS:
        matches = list(RESOURCES.glob(prefix + '*.pdf'))
        if len(matches) != 1:
            raise RuntimeError(f'{ident}: expected one match, got {matches}')
        items.append((ident, matches[0]))
    for i, path in enumerate(sorted((ROOT / 'documents').iterdir()), 1):
        if path.is_file():
            items.append((f'D{i:02d}', path))
    manifest = []
    for ident, path in items:
        if args.only != 'all' and ident not in args.only.split(','):
            continue
        out = WORK / '提取' / ident
        suffix = path.suffix.lower()
        try:
            if suffix == '.pdf':
                meta = extract_pdf(path, out)
            elif suffix == '.docx':
                meta = extract_docx(path, out)
            elif suffix == '.pptx':
                meta = extract_pptx(path, out)
            elif suffix == '.xlsx':
                meta = extract_xlsx(path, out)
            elif suffix == '.md':
                dump(out / '全文.md', path.read_text(encoding='utf-8'))
                meta = {'kind': 'markdown'}
            else:
                meta = {'error': 'unsupported format'}
            record = {'id': ident, 'source': str(path), 'extracted': str(out),
                      'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), **meta}
        except Exception as exc:
            record = {'id': ident, 'source': str(path), 'error': type(exc).__name__ + ': ' + str(exc.reason if isinstance(exc, UnicodeError) else exc)[:300]}
        manifest.append(record)
        print(json.dumps(record, ensure_ascii=False), flush=True)
    dump(WORK / ('清单.json' if args.only == 'all' else '清单_' + args.only.replace(',', '_') + '.json'),
         json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
